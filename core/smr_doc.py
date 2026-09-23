"""Suspicious Matter Report (SMR) draft as a Word document.

Laid out after the sections an SMR asks for in AUSTRAC Online (reporting entity, the person or organisation
the suspicion relates to, the designated service and transactions, the grounds for suspicion, and the action
taken), so the MLRO can review the draft and key it into AUSTRAC Online. This is a draft for review, not a
lodged report, and every figure in it is synthetic.
"""
import io

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from . import case as C
from .ref import FIRM_NAME, MLRO_NAME, USER_NAME, USER_ROLE, TYPOLOGIES

INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x67, 0x70)
ACCENT = RGBColor(0x5C, 0x6E, 0x10)      # deep olive — the app's lime, darkened for print
RED = RGBColor(0xB4, 0x2D, 0x24)
HEAD_FILL = "E8EDD2"
LABEL_FILL = "F3F4F1"


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _bottom_rule(paragraph, color="5C6E10", size=8):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size)); b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    bdr.append(b); pPr.append(bdr)


def _run(p, text, bold=False, size=10, color=INK, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size, r.font.color.rgb = Pt(size), color
    return r


def _heading(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12), Pt(4)
    _run(p, f"{number}  ", bold=True, size=12, color=ACCENT)
    _run(p, text, bold=True, size=12)
    _bottom_rule(p, size=4)
    return p


def _kv_table(doc, rows, widths=(5.2, 11.3)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(widths[0]), Cm(widths[1])
        _shade(cells[0], LABEL_FILL)
        cells[0].paragraphs[0].text = ""
        _run(cells[0].paragraphs[0], k, bold=True, size=9, color=MUTED)
        _run(cells[1].paragraphs[0], str(v), size=9.5)
    return t


def _grid(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i]); _shade(c, HEAD_FILL)
        _run(c.paragraphs[0], h, bold=True, size=9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            _run(p, str(v), size=9)
            if isinstance(v, str) and v.startswith("AUD"):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return t


def _bullets(doc, items, color=INK):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        _run(p, it, size=10, color=color)


def actions_taken(c, res, escalated_on=None):
    """Section 6 in the past tense — what the firm has done, as an SMR records it."""
    out = []
    if c.get("alert_id"):
        out.append(f"Alert {c['alert_id']} raised by the transaction-monitoring rule '{c.get('typology')}' and reviewed by the analyst.")
    if escalated_on is not None:
        out.append(f"Escalated to the MLRO ({MLRO_NAME}) on {pd.Timestamp(escalated_on):%d %B %Y}.")
    if "sanctions_confirmed" in c.get("hard", []):
        out.append("Instructions suspended and funds held; the Australian Sanctions Office (DFAT) to be notified.")
    elif res["escalate"]:
        out.append("Matter placed on hold; no trust money released pending the MLRO's decision.")
    missing = [m for m in C.MITIGANTS if m not in c.get("mitigants", [])]
    if missing:
        out.append("Enhanced customer due diligence requested: " + ", ".join(C.MITIGANTS[m][0][0].lower() + C.MITIGANTS[m][0][1:] for m in missing) + ".")
    if res["ttr"]:
        out.append("Threshold transaction report prepared for the physical cash component.")
    out.append("Client not informed of the review or of any report (tipping-off prohibition, s123).")
    out.append("Case file and transaction records retained for 7 years.")
    return out


def build_smr_docx(c, today, smr_ref=None, rationale=None, escalated_on=None, analyst=None):
    """Return the SMR draft as .docx bytes for case dict `c` (see core.case)."""
    res = C.assess(c, today)
    today = pd.Timestamp(today)
    tf = "terrorism" in c.get("hard", [])
    due = today + pd.Timedelta(hours=24) if tf else C.add_business_days(today, 3)
    smr_ref = smr_ref or f"SMR-DRAFT-{(c.get('alert_id') or 'MANUAL').replace('ALT-', '')}"

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10)
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # header / footer
    hp = sec.header.paragraphs[0]
    _run(hp, "CONFIDENTIAL — DRAFT FOR MLRO REVIEW", bold=True, size=8, color=RED)
    _run(hp, f"   ·   {FIRM_NAME} · AML/CTF Compliance", size=8, color=MUTED)
    fp = sec.footer.paragraphs[0]
    _run(fp, "Tipping off is prohibited (AML/CTF Act s123): do not disclose to the client or any third party that an SMR "
             "has been or may be lodged. Synthetic data generated for a portfolio project — not a real report.", size=7.5, color=MUTED, italic=True)

    # title block
    p = doc.add_paragraph()
    _run(p, "Suspicious Matter Report", bold=True, size=20)
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    _run(p, "Draft prepared for lodgement through AUSTRAC Online", size=10.5, color=MUTED)
    _bottom_rule(p, size=12)

    status_color = RED if res["smr_recommended"] or tf else ACCENT
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    _run(p, "Recommendation:  ", bold=True, size=10.5)
    _run(p, f"{res['title']}  ·  score {res['score']:.0f}/100", bold=True, size=10.5, color=status_color)

    # 1. report details
    _heading(doc, "1", "Report details")
    _kv_table(doc, [
        ("Internal reference", smr_ref),
        ("Source alert", c.get("alert_id") or "Manual case entry"),
        ("Suspected offence", "Terrorism financing" if tf else ("Sanctions breach / money laundering" if "sanctions_confirmed" in c.get("hard", [])
                                                               else f"Money laundering — {c.get('typology', 'see grounds')}")),
        ("Escalated to MLRO", f"{pd.Timestamp(escalated_on):%d %B %Y}" if escalated_on is not None else "Not yet escalated"),
        ("Suspicion formed / draft date", f"{today:%d %B %Y}"),
        ("Lodgement due", f"{due:%d %B %Y}" + (" (24 hours — terrorism financing)" if tf else " (3 business days from suspicion)")),
        ("Reporting entity", f"{FIRM_NAME} — {c.get('branch', 'Richmond')} office"),
        ("Designated service", c["product"]),
    ])

    # 2. subject
    _heading(doc, "2", "Person or organisation the suspicion relates to")
    _kv_table(doc, [
        ("Name", c.get("client") or "—"),
        ("Client ID", c.get("client_id") or "—"),
        ("Client type", c["client_type"]),
        ("State / territory", c.get("state", "—")),
        ("Delivery channel", c["channel"]),
        ("Foreign dimension", c["foreign_tier"]),
        ("Firm risk rating", f"{res['client_tier']} ({res['client_score']:.0f}/100, AUSTRAC 4-factor model)"),
    ])
    owners = C.beneficial_owners(c.get("client"))
    if owners:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
        _run(p, "Beneficial ownership on file", bold=True, size=10)
        _grid(doc, ["Owner", "Interest"], [(o, f"{pct}%") for o, pct in owners], widths=(12.5, 4.0))

    # 3. transactions
    _heading(doc, "3", "Transactions")
    tx = c.get("transactions") or []
    if tx:
        _grid(doc, ["Date", "Transaction", "Direction", "Counterparty", "Amount"],
              [(f"{d:%d %b %Y}", kind, direction, cp, f"AUD {amt:,.0f}") for d, kind, amt, direction, cp in tx],
              widths=(2.4, 5.2, 1.9, 3.6, 3.4))
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        _run(p, f"Total: AUD {sum(t[2] for t in tx):,.0f} across {len(tx)} transaction(s).", bold=True, size=9.5)
    else:
        p = doc.add_paragraph()
        _run(p, f"AUD {c.get('amount', 0):,.0f} across {c.get('n_txns', 1)} transaction(s) — itemise before lodgement.", size=10)
    if res["ttr"]:
        p = doc.add_paragraph()
        _run(p, "Physical cash of AUD 10,000 or more was received: a threshold transaction report (TTR) is also due within "
                "10 business days, separately from this SMR.", size=9.5, color=RED)

    # 4. grounds
    _heading(doc, "4", "Grounds for suspicion")
    p = doc.add_paragraph()
    _run(p, "First draft generated from the case facts — the analyst and MLRO must review and edit before lodgement.", size=8.5, color=MUTED, italic=True)
    for para in C.grounds_narrative(c, res):
        q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(6)
        _run(q, para, size=10.5)
    if rationale:
        q = doc.add_paragraph()
        _run(q, "Analyst's note: ", bold=True, size=10.5)
        _run(q, rationale, size=10.5)
    typ = TYPOLOGIES.get(c.get("typology"))
    if typ:
        q = doc.add_paragraph()
        _run(q, "Monitoring rule triggered: ", bold=True, size=10)
        _run(q, f"{c['typology']} — {typ['desc']}", size=10, color=MUTED)

    # 5. indicators
    _heading(doc, "5", "Indicators and verification")
    flags = [C.RED_FLAGS[f][0] for f in c.get("flags", []) if f in C.RED_FLAGS] + [C.HARD_TRIGGERS[h] for h in c.get("hard", []) if h in C.HARD_TRIGGERS]
    p = doc.add_paragraph(); _run(p, "Red flags observed", bold=True, size=10)
    _bullets(doc, flags or ["None recorded"])
    p = doc.add_paragraph(); _run(p, "Verification obtained", bold=True, size=10)
    got = [C.MITIGANTS[m][0] for m in c.get("mitigants", []) if m in C.MITIGANTS]
    _bullets(doc, got or ["None — no mitigating evidence obtained"])
    missing = [C.MITIGANTS[m][0] for m in C.MITIGANTS if m not in c.get("mitigants", [])]
    if missing:
        p = doc.add_paragraph(); _run(p, "Outstanding", bold=True, size=10)
        _bullets(doc, missing, color=MUTED)

    # 6. action taken
    _heading(doc, "6", "Action taken by the reporting entity")
    _bullets(doc, actions_taken(c, res, escalated_on))

    # 7. sign-off
    _heading(doc, "7", "Review and sign-off")
    t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; t.autofit = False
    for i, (role, name) in enumerate([(f"Prepared by — {USER_ROLE}", analyst or USER_NAME), ("Reviewed by — MLRO", MLRO_NAME)]):
        cell = t.rows[0].cells[i]; cell.width = Cm(8.25)
        _run(cell.paragraphs[0], role, bold=True, size=9, color=MUTED)
        for line in (name, "Signature: ______________________", "Date: ____ / ____ / ________"):
            q = cell.add_paragraph(); _run(q, line, size=10)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    _run(p, "MLRO decision:  ☐ Suspicion formed — lodge SMR     ☐ Return for further information     ☐ No suspicion — close", size=10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

"""
report_utils.py
Automated detection of common ML/TF typologies in a customer's transaction
history, and generation of an AUSTRAC-aligned Suspicious Matter Report (SMR)
narrative.

AUSTRAC expects an effective SMR grounds-for-suspicion (GFS) narrative to
cover six elements, in standard case (not ALL CAPS), written as clear,
factual, specific prose organised under logical headings:
    WHO   - full details of the person/entity and relationship to the report
    WHAT  - what happened (transaction/service/behaviour, amounts, products)
    WHEN  - when the activity was first observed and over what period
    WHERE - where the activity occurred (channel, jurisdiction)
    HOW   - how the activity was carried out / the mechanism
    WHY   - why the reporting entity believes it is suspicious

This module builds that narrative programmatically from the demo (or live)
transaction data, then renders it as plain text and as a formatted .docx.
"""

from __future__ import annotations
from dataclasses import dataclass, field
from datetime import datetime
import io

import pandas as pd

REPORTING_THRESHOLD = 10000  # AUD - IFTI/TTR structuring threshold reference


# --------------------------------------------------------------------------
# TYPOLOGY DETECTION
# --------------------------------------------------------------------------
@dataclass
class Finding:
    typology: str
    summary: str
    evidence: pd.DataFrame
    severity: str  # Low / Medium / High / Critical


def detect_structuring(txns: pd.DataFrame) -> Finding | None:
    cash = txns[(txns["channel"] == "Cash Deposit") & (txns["amount"] >= 8500) & (txns["amount"] < REPORTING_THRESHOLD)]
    if cash.empty:
        return None
    cash = cash.sort_values("txn_date")
    cash["window"] = (cash["txn_date"] - cash["txn_date"].min()).dt.days // 10
    for _, grp in cash.groupby("window"):
        if len(grp) >= 3:
            total = grp["amount"].sum()
            span_days = (grp["txn_date"].max() - grp["txn_date"].min()).days
            summary = (
                f"{len(grp)} cash deposits totalling ${total:,.2f} were made over "
                f"{max(span_days,1)} day(s), each individually below the ${REPORTING_THRESHOLD:,.0f} "
                "threshold at which a threshold transaction report would ordinarily apply."
            )
            return Finding("Structuring (Smurfing)", summary, grp, "High")
    return None


def detect_rapid_movement(txns: pd.DataFrame) -> Finding | None:
    inbound = txns[(txns["direction"] == "Inbound") & (txns["amount"] >= 25000)].sort_values("txn_date")
    for _, in_row in inbound.iterrows():
        window_end = in_row["txn_date"] + pd.Timedelta(days=3)
        out = txns[
            (txns["direction"] == "Outbound")
            & (txns["txn_date"] >= in_row["txn_date"])
            & (txns["txn_date"] <= window_end)
        ]
        if not out.empty and out["amount"].sum() >= 0.6 * in_row["amount"]:
            evidence = pd.concat([in_row.to_frame().T, out])
            hours_elapsed = max((out["txn_date"].max() - in_row["txn_date"]).total_seconds() / 3600, 1)
            elapsed_txt = (
                f"{hours_elapsed:.0f} hour(s)" if hours_elapsed < 24
                else f"{hours_elapsed/24:.0f} day(s)"
            )
            summary = (
                f"An inbound credit of ${in_row['amount']:,.2f} received via {in_row['channel'].lower()} "
                f"on {in_row['txn_date'].date()} was substantially depleted within {elapsed_txt} via "
                f"{out.shape[0]} outbound transaction(s) totalling ${out['amount'].sum():,.2f}, consistent "
                "with a pass-through / layering pattern rather than legitimate account usage."
            )
            return Finding("Rapid Movement of Funds", summary, evidence, "High")
    return None


def detect_high_risk_wires(txns: pd.DataFrame, high_risk_countries: set[str]) -> Finding | None:
    wires = txns[
        (txns["channel"] == "International Wire")
        & (txns["counterparty_country"].isin(high_risk_countries))
    ]
    if len(wires) < 2:
        return None
    total = wires["amount"].sum()
    countries = ", ".join(sorted(wires["counterparty_country"].unique()))
    summary = (
        f"{len(wires)} international wire transfers totalling ${total:,.2f} were sent to or received from "
        f"counterparties in higher-risk jurisdictions ({countries}), a corridor associated with elevated "
        "money laundering and trade-based value transfer risk."
    )
    return Finding("High-Risk Corridor Wire Activity", summary, wires, "Medium")


HIGH_RISK_COUNTRIES_DEFAULT = {"Myanmar", "Cambodia", "Panama", "Cyprus", "Nigeria", "Papua New Guinea"}


def run_all_detections(txns: pd.DataFrame) -> list[Finding]:
    findings = []
    for fn in (
        lambda t: detect_structuring(t),
        lambda t: detect_rapid_movement(t),
        lambda t: detect_high_risk_wires(t, HIGH_RISK_COUNTRIES_DEFAULT),
    ):
        result = fn(txns)
        if result is not None:
            findings.append(result)
    return findings


# --------------------------------------------------------------------------
# SMR NARRATIVE ASSEMBLY (AUSTRAC-aligned)
# --------------------------------------------------------------------------
@dataclass
class SMRReport:
    customer: dict
    findings: list[Finding]
    screening_hits: pd.DataFrame
    reference: str
    generated_at: datetime = field(default_factory=datetime.now)

    # -------- section builders --------
    def part_a_reporting_entity(self) -> str:
        return (
            "Part A - Reporting Entity Details\n"
            "----------------------------------\n"
            "Reporting entity: [Enter registered business name]\n"
            "AUSTRAC reference number: [Enter reporting entity AUSTRAC ID]\n"
            f"Report reference: {self.reference}\n"
            f"Report prepared: {self.generated_at.strftime('%d %B %Y, %H:%M')} (AEST)\n"
            "Reason for report: Suspicion of money laundering under the AML/CTF Act 2006 (Cth)\n"
        )

    def part_c_suspicious_person(self) -> str:
        c = self.customer
        return (
            "Part C - Details of Suspicious Person / Entity\n"
            "-----------------------------------------------\n"
            f"Customer reference: {c['customer_id']}\n"
            f"Full name / registered name: {c['name']}\n"
            f"Customer type: {c['customer_type']}\n"
            f"Occupation / industry: {c.get('occupation', 'N/A')} / {c['industry']}\n"
            f"Country of residence / incorporation: {c['country']}\n"
            f"Relationship to reporting entity: Existing customer, onboarded "
            f"{pd.Timestamp(c['onboarding_date']).strftime('%d %B %Y')}\n"
            f"Current internal risk rating: {c['risk_level']} ({c['risk_score']}/100)\n"
        )

    def part_b_transaction_summary(self) -> str:
        lines = ["Part B - Transaction / Activity Summary", "-----------------------------------------"]
        if not self.findings:
            lines.append("No specific transaction pattern met automated detection thresholds. "
                         "Manual grounds for suspicion should be entered below.")
            return "\n".join(lines) + "\n"
        for f in self.findings:
            lines.append(f"* {f.typology} (severity: {f.severity}): {f.summary}")
        return "\n".join(lines) + "\n"

    def grounds_for_suspicion(self) -> str:
        c = self.customer
        article = "an" if c['customer_type'].lower().startswith(("i", "a", "e", "o", "u")) else "a"
        who = (
            f"{c['name']} ({c['customer_id']}) is {article} {c['customer_type'].lower()} customer of the reporting "
            f"entity, resident/incorporated in {c['country']}, operating in the {c['industry']} sector. "
            f"The customer holds an internal risk rating of {c['risk_level']}."
        )

        if self.findings:
            what_parts = [f.summary for f in self.findings]
            what = " ".join(what_parts)
            when = self._when_range()
            where = self._where_summary()
            how = self._how_summary()
        else:
            what = "[Describe the specific transaction(s) or behaviour that gave rise to suspicion.]"
            when = "[Insert date(s) or date range over which the activity occurred.]"
            where = "[Insert channel/branch/jurisdiction where the activity took place.]"
            how = "[Describe the mechanism - e.g. cash, wire, third party, digital currency.]"

        why_reasons = [
            "the pattern of activity is inconsistent with the stated purpose and expected transaction "
            "profile of the account established at onboarding",
        ]
        if any(f.typology.startswith("Structuring") for f in self.findings):
            why_reasons.append(
                "deposit amounts were structured to remain below the $10,000 threshold at which a "
                "threshold transaction report would ordinarily be triggered, which is a recognised "
                "indicator of structuring/smurfing under s.142 of the AML/CTF Act"
            )
        if any(f.typology.startswith("Rapid") for f in self.findings):
            why_reasons.append(
                "funds were moved out of the account shortly after receipt, with no apparent business or "
                "personal rationale, consistent with layering to obscure the origin of funds"
            )
        if any("High-Risk" in f.typology for f in self.findings):
            why_reasons.append(
                "counterparties are located in jurisdictions assessed by the reporting entity as higher risk "
                "for money laundering, terrorism financing or corruption"
            )
        if not self.screening_hits.empty:
            open_hits = self.screening_hits[self.screening_hits["status"].isin(["Open", "Escalated", "Under Review"])]
            if not open_hits.empty:
                types = ", ".join(sorted(open_hits["match_type"].unique()))
                why_reasons.append(
                    f"the customer has unresolved watchlist screening result(s) of type: {types}"
                )
        why = "Considered together, " + "; and ".join(why_reasons) + "."

        narrative = (
            "Grounds for Suspicion\n"
            "----------------------\n\n"
            f"Who\n{who}\n\n"
            f"What\n{what}\n\n"
            f"When\n{when}\n\n"
            f"Where\n{where}\n\n"
            f"How\n{how}\n\n"
            f"Why\n{why}\n"
        )
        return narrative

    def _when_range(self) -> str:
        all_ev = pd.concat([f.evidence for f in self.findings])
        start = pd.to_datetime(all_ev["txn_date"]).min()
        end = pd.to_datetime(all_ev["txn_date"]).max()
        if start.date() == end.date():
            return f"The activity was first observed on {start.strftime('%d %B %Y')}."
        return (
            f"The activity was observed between {start.strftime('%d %B %Y')} and "
            f"{end.strftime('%d %B %Y')}, a period of {(end - start).days} day(s)."
        )

    def _where_summary(self) -> str:
        all_ev = pd.concat([f.evidence for f in self.findings])
        channels = ", ".join(sorted(all_ev["channel"].unique()))
        countries = sorted(set(all_ev.get("counterparty_country", pd.Series(dtype=str)).dropna().unique()))
        country_txt = f" involving counterparties in {', '.join(countries)}" if countries else ""
        return f"The activity took place via {channels}{country_txt}."

    def _how_summary(self) -> str:
        mechanisms = []
        for f in self.findings:
            if f.typology.startswith("Structuring"):
                mechanisms.append("multiple cash deposits kept just under the reporting threshold")
            elif f.typology.startswith("Rapid"):
                mechanisms.append("prompt transfer of incoming funds out of the account")
            elif "High-Risk" in f.typology:
                mechanisms.append("international wire transfers routed through higher-risk corridors")
        return ("The suspected activity was carried out through " + "; and ".join(mechanisms) + ".") if mechanisms else \
            "[Describe the mechanism used to carry out the activity.]"

    def full_text(self) -> str:
        sections = [
            "SUSPICIOUS MATTER REPORT (SMR) - DRAFT",
            "Prepared for internal review prior to AUSTRAC Online lodgement",
            "=" * 70,
            "",
            self.part_a_reporting_entity(),
            self.part_c_suspicious_person(),
            self.part_b_transaction_summary(),
            self.grounds_for_suspicion(),
            "Statutory Timeframe\n--------------------",
            ("This report relates to a suspected money laundering / structuring matter and should be "
             "lodged with AUSTRAC within 3 business days of the suspicion being formed. If any element "
             "of the activity is suspected to relate to terrorism financing, the report must be lodged "
             "within 24 hours.\n"),
            "Reviewing Officer Certification\n---------------------------------",
            ("Prepared by: [Compliance analyst name]        Reviewed by: [MLRO / delegate name]\n"
             "I have reviewed the customer's transaction history, KYC file and screening results and "
             "consider that reasonable grounds for suspicion exist as described above.\n"
             "Signature: ______________________        Date: ______________________\n"),
            "-" * 70,
            "Note: This draft was generated automatically from transaction monitoring data to accelerate "
            "case review. All bracketed [ ] fields, entity details and the grounds-for-suspicion narrative "
            "must be verified, edited and approved by a qualified compliance officer before lodgement with "
            "AUSTRAC. Do not submit unedited system output.",
        ]
        return "\n".join(sections)


def build_smr(customer_row: pd.Series, txns: pd.DataFrame, screening_hits: pd.DataFrame) -> SMRReport:
    cust_txns = txns[txns["customer_id"] == customer_row["customer_id"]].copy()
    findings = run_all_detections(cust_txns)
    cust_screen = screening_hits[screening_hits["customer_id"] == customer_row["customer_id"]].copy()
    reference = f"SMR-{customer_row['customer_id']}-{datetime.now().strftime('%Y%m%d%H%M')}"
    return SMRReport(
        customer=customer_row.to_dict(),
        findings=findings,
        screening_hits=cust_screen,
        reference=reference,
    )


# --------------------------------------------------------------------------
# DOCX EXPORT
# --------------------------------------------------------------------------
def smr_to_docx_bytes(report: SMRReport) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    title = doc.add_heading("Suspicious Matter Report (SMR) - Draft", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Prepared for internal compliance review prior to AUSTRAC Online lodgement")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    warn = doc.add_paragraph()
    run = warn.add_run(
        "DRAFT ONLY - all bracketed fields and the grounds-for-suspicion narrative must be reviewed, "
        "verified and approved by a qualified compliance officer / MLRO before lodgement with AUSTRAC."
    )
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)

    def add_section(heading: str, body: str):
        doc.add_heading(heading, level=2)
        for para in body.strip("\n").split("\n"):
            if para.strip() == "":
                continue
            p = doc.add_paragraph(para)
            p.paragraph_format.space_after = Pt(4)

    add_section("Part A - Reporting Entity Details", report.part_a_reporting_entity())
    add_section("Part C - Details of Suspicious Person / Entity", report.part_c_suspicious_person())
    add_section("Part B - Transaction / Activity Summary", report.part_b_transaction_summary())

    doc.add_heading("Grounds for Suspicion", level=2)
    gfs = report.grounds_for_suspicion().split("Grounds for Suspicion\n----------------------\n\n")[-1]
    for block in gfs.strip().split("\n\n"):
        lines = block.split("\n", 1)
        if len(lines) == 2:
            heading, text_body = lines
            hp = doc.add_paragraph()
            hp.add_run(heading).bold = True
            doc.add_paragraph(text_body)

    add_section(
        "Statutory Timeframe",
        "This report relates to a suspected money laundering / structuring matter and should be lodged "
        "with AUSTRAC within 3 business days of the suspicion being formed. If any element of the activity "
        "is suspected to relate to terrorism financing, the report must be lodged within 24 hours.",
    )
    add_section(
        "Reviewing Officer Certification",
        "Prepared by: [Compliance analyst name]      Reviewed by: [MLRO / delegate name]\n"
        "I have reviewed the customer's transaction history, KYC file and screening results and consider "
        "that reasonable grounds for suspicion exist as described above.\n"
        "Signature: ______________________      Date: ______________________",
    )

    if report.findings:
        doc.add_heading("Supporting Evidence - Flagged Transactions", level=2)
        for f in report.findings:
            doc.add_paragraph(f"{f.typology} ({f.severity} severity)", style="Intense Quote")
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, col in enumerate(["Txn ID", "Date", "Amount", "Direction", "Channel"]):
                hdr[i].text = col
            for _, row in f.evidence.head(8).iterrows():
                cells = table.add_row().cells
                cells[0].text = str(row.get("txn_id", ""))
                cells[1].text = str(pd.Timestamp(row["txn_date"]).date())
                cells[2].text = f"${row['amount']:,.2f}"
                cells[3].text = str(row.get("direction", ""))
                cells[4].text = str(row.get("channel", ""))
            doc.add_paragraph()

    footer = doc.add_paragraph(f"Generated {report.generated_at.strftime('%d %B %Y %H:%M')} | Reference {report.reference}")
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

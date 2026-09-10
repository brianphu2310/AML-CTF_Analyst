"""
onboarding_utils.py
Rule-based CDD/EDD scoring engine and Customer Due Diligence (CDD) record
generator used by the "New Client Onboarding" page.

The scoring model is intentionally transparent (a running list of weighted
risk factors) so a compliance officer can see exactly why a score landed
where it did - this mirrors how AML/CTF Programs typically document their
customer risk assessment methodology (customer type, geography, industry/
occupation, delivery channel, products/services, and PEP/adverse-media
status are the standard risk categories under an AML/CTF Program).
"""

from __future__ import annotations
from dataclasses import dataclass, field
from datetime import datetime
import io

import pandas as pd

HIGH_RISK_COUNTRIES = {"Myanmar", "Cambodia", "Panama", "Cyprus", "Nigeria", "Papua New Guinea", "Iran", "North Korea"}
HIGH_RISK_INDUSTRIES = {
    "Remittance/MSB", "Precious Metals", "Digital Currency Exchange", "Used Vehicle Sales",
    "Real Estate", "Import/Export", "Gambling/Wagering", "Cash-Intensive Retail",
}
HIGH_RISK_PURPOSES = {"Remittance", "Trade Finance", "Property Purchase", "Investment - Offshore"}
HIGH_RISK_SOF = {"Loan (private/informal)", "Gift", "Gambling winnings", "Other / Unspecified"}


@dataclass
class RiskFactor:
    label: str
    points: int
    note: str = ""


@dataclass
class OnboardingResult:
    factors: list[RiskFactor]
    score: int
    risk_level: str
    cdd_tier: str
    decision: str
    screening_hits: list[dict]

    def factor_table(self) -> pd.DataFrame:
        return pd.DataFrame([{"Factor": f.label, "Points": f.points, "Note": f.note} for f in self.factors])


def _name_hash_signal(name: str) -> int:
    """Deterministic pseudo-signal so the same name always screens the same
    way in this demo (stands in for a real watchlist-matching engine)."""
    return sum(ord(c) for c in name.lower()) if name else 0


def run_screening(full_name: str, is_business: bool, self_declared_pep: bool) -> list[dict]:
    """Simulated PEP/Sanctions/Adverse Media screening. In production this
    calls a real watchlist provider (e.g. Dow Jones, World-Check, Refinitiv)."""
    hits = []
    signal = _name_hash_signal(full_name)

    if self_declared_pep:
        hits.append({
            "type": "PEP",
            "detail": "Self-declared politically exposed person",
            "score": 100,
            "source": "Onboarding declaration",
        })

    # Simulated fuzzy watchlist signal - deterministic but looks probabilistic
    if signal % 17 == 0 and full_name:
        hits.append({
            "type": "Sanctions",
            "detail": "Potential name match on consolidated sanctions list - requires manual review",
            "score": 78,
            "source": "DFAT Consolidated List (simulated)",
        })
    if signal % 11 == 0 and full_name:
        hits.append({
            "type": "Adverse Media",
            "detail": "Potential adverse media match - unverified, requires analyst review",
            "score": 62,
            "source": "News archive screening (simulated)",
        })
    if signal % 23 == 0 and full_name and not is_business:
        hits.append({
            "type": "PEP",
            "detail": "Potential match to foreign PEP database entry - requires manual review",
            "score": 71,
            "source": "World-Check (simulated)",
        })
    return hits


def score_onboarding(answers: dict, screening_hits: list[dict]) -> OnboardingResult:
    factors: list[RiskFactor] = []

    # --- Customer type & geography ---
    if answers.get("customer_type") == "Business":
        factors.append(RiskFactor("Business customer", 5, "Corporate structures carry inherently higher complexity"))
        if answers.get("structure_type") in ("Trust", "Foreign Subsidiary"):
            factors.append(RiskFactor(f"Structure type: {answers['structure_type']}", 12,
                                       "Trusts/foreign subsidiaries can obscure beneficial ownership"))
        ubos = answers.get("ubos", [])
        if any(u.get("is_pep") for u in ubos):
            factors.append(RiskFactor("PEP among beneficial owners", 30, "Enhanced due diligence required"))
        if any(u.get("ownership_pct", 0) and u.get("nationality") in HIGH_RISK_COUNTRIES for u in ubos):
            factors.append(RiskFactor("Beneficial owner in high-risk jurisdiction", 15))
    else:
        factors.append(RiskFactor("Individual customer", 0))

    country = answers.get("country", "")
    if country in HIGH_RISK_COUNTRIES:
        factors.append(RiskFactor(f"Residence/incorporation: {country}", 25, "Assessed as a higher-risk jurisdiction"))
    elif country and country != "Australia":
        factors.append(RiskFactor(f"Residence/incorporation: {country}", 8, "Foreign, standard risk jurisdiction"))

    # --- Industry / occupation ---
    industry = answers.get("industry", "")
    if industry in HIGH_RISK_INDUSTRIES:
        factors.append(RiskFactor(f"Industry: {industry}", 20, "Cash-intensive / typology-prone sector"))

    if answers.get("cash_intensive"):
        factors.append(RiskFactor("Self-identified as cash-intensive business", 15))

    # --- Purpose / source of funds ---
    purpose = answers.get("purpose", "")
    if purpose in HIGH_RISK_PURPOSES:
        factors.append(RiskFactor(f"Purpose of relationship: {purpose}", 10))

    sof = answers.get("source_of_funds", "")
    if sof in HIGH_RISK_SOF:
        factors.append(RiskFactor(f"Source of funds: {sof}", 15, "Less verifiable source"))

    # --- Expected activity ---
    volume = answers.get("expected_volume", "")
    if volume in ("$50,000 - $250,000 / month", "Over $250,000 / month"):
        factors.append(RiskFactor(f"Expected transaction volume: {volume}", 12))

    channels = answers.get("expected_channels", [])
    if "Cash" in channels:
        factors.append(RiskFactor("Expected use of cash", 10))
    if "Digital Currency" in channels:
        factors.append(RiskFactor("Expected use of digital currency", 15))

    countries_txn = answers.get("expected_countries", [])
    hr_overlap = [c for c in countries_txn if c in HIGH_RISK_COUNTRIES]
    if hr_overlap:
        factors.append(RiskFactor(f"Expected transacting with: {', '.join(hr_overlap)}", 18,
                                   "Higher-risk counterparty jurisdiction(s)"))

    # --- Identity verification ---
    if answers.get("verification_method") == "Non face-to-face (digital)":
        factors.append(RiskFactor("Non face-to-face onboarding", 8, "Elevated impersonation/fraud risk"))

    # --- Screening ---
    for hit in screening_hits:
        pts = {"Sanctions": 60, "PEP": 30, "Adverse Media": 20}.get(hit["type"], 15)
        factors.append(RiskFactor(f"Screening hit: {hit['type']} - {hit['detail']}", pts))

    score = max(0, min(100, sum(f.points for f in factors)))

    if score >= 75 or any(h["type"] == "Sanctions" for h in screening_hits):
        risk_level = "Critical"
    elif score >= 50:
        risk_level = "High"
    elif score >= 25:
        risk_level = "Medium"
    else:
        risk_level = "Low"

    cdd_tier = {
        "Low": "Simplified Due Diligence (SDD)",
        "Medium": "Standard Customer Due Diligence (CDD)",
        "High": "Enhanced Due Diligence (EDD)",
        "Critical": "Enhanced Due Diligence (EDD) + Senior Management Sign-off",
    }[risk_level]

    if any(h["type"] == "Sanctions" for h in screening_hits):
        decision = "DO NOT ONBOARD - Refer immediately to MLRO for sanctions review before any service is provided"
    elif risk_level == "Critical":
        decision = "Refer to MLRO for approval - EDD must be completed before account activation"
    elif risk_level == "High":
        decision = "Proceed with Enhanced Due Diligence - senior compliance approval required"
    elif risk_level == "Medium":
        decision = "Proceed with standard onboarding - periodic review in 12 months"
    else:
        decision = "Proceed with simplified onboarding - periodic review in 24 months"

    return OnboardingResult(factors, score, risk_level, cdd_tier, decision, screening_hits)


# --------------------------------------------------------------------------
# CDD RECORD TEXT / DOCX
# --------------------------------------------------------------------------
def cdd_record_text(answers: dict, result: OnboardingResult) -> str:
    lines = [
        "CUSTOMER DUE DILIGENCE (CDD) RECORD",
        "New Client Onboarding - Compliance File",
        "=" * 70,
        f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}",
        "",
        "1. Customer Identification",
        "-" * 30,
        f"Customer type: {answers.get('customer_type')}",
        f"Full name / legal name: {answers.get('full_name')}",
        f"{'Date of birth' if answers.get('customer_type')=='Individual' else 'Incorporation date'}: "
        f"{answers.get('dob_or_incorp')}",
        f"{'Nationality' if answers.get('customer_type')=='Individual' else 'Structure type'}: "
        f"{answers.get('nationality') if answers.get('customer_type')=='Individual' else answers.get('structure_type')}",
        f"Country of residence / incorporation: {answers.get('country')}",
        f"Occupation / industry: {answers.get('occupation') or answers.get('industry')}",
        f"Identification document: {answers.get('id_type')} {answers.get('id_number')}",
        f"Verification method: {answers.get('verification_method')}",
        "",
    ]

    if answers.get("customer_type") == "Business":
        lines += ["2. Beneficial Ownership", "-" * 30]
        ubos = answers.get("ubos", [])
        if ubos:
            for u in ubos:
                pep = " [PEP]" if u.get("is_pep") else ""
                lines.append(f"- {u.get('name')} ({u.get('role')}) - {u.get('ownership_pct')}% - "
                             f"{u.get('nationality')}{pep}")
        else:
            lines.append("No beneficial owners recorded.")
        lines.append("")

    lines += [
        "3. Relationship & Activity Profile",
        "-" * 30,
        f"Purpose of relationship: {answers.get('purpose')}",
        f"Source of funds: {answers.get('source_of_funds')}",
        f"Source of wealth: {answers.get('source_of_wealth')}",
        f"Expected transaction volume: {answers.get('expected_volume')}",
        f"Expected channels: {', '.join(answers.get('expected_channels', []))}",
        f"Expected transacting countries: {', '.join(answers.get('expected_countries', []))}",
        f"Cash-intensive business: {'Yes' if answers.get('cash_intensive') else 'No'}",
        "",
        "4. Screening Results",
        "-" * 30,
    ]
    if result.screening_hits:
        for h in result.screening_hits:
            lines.append(f"- {h['type']}: {h['detail']} (score {h['score']}, source: {h['source']})")
    else:
        lines.append("No PEP, sanctions or adverse media matches identified.")

    lines += [
        "",
        "5. Risk Assessment",
        "-" * 30,
    ]
    for f in result.factors:
        note = f" - {f.note}" if f.note else ""
        lines.append(f"- {f.label}: +{f.points}{note}")
    lines += [
        f"\nTotal risk score: {result.score}/100",
        f"Risk rating: {result.risk_level}",
        f"CDD tier required: {result.cdd_tier}",
        f"Onboarding decision: {result.decision}",
        "",
        "6. Compliance Officer Sign-off",
        "-" * 30,
        "Prepared by: [Compliance analyst name]        Reviewed by: [MLRO / delegate name]",
        "I confirm the identification, verification and risk assessment steps above have been completed "
        "in accordance with the AML/CTF Program.",
        "Signature: ______________________        Date: ______________________",
        "",
        "-" * 70,
        "Note: This record was generated to accelerate onboarding documentation. All fields must be "
        "verified against original/certified identification documents, and any Critical/High risk rating "
        "or screening hit must be reviewed and approved by the MLRO before the account is activated.",
    ]
    return "\n".join(str(l) for l in lines)


def cdd_record_to_docx_bytes(answers: dict, result: OnboardingResult) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading("Customer Due Diligence (CDD) Record", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("New Client Onboarding - Compliance File")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    risk_color = {
        "Low": RGBColor(0x1F, 0x6F, 0x50), "Medium": RGBColor(0x8A, 0x6D, 0x1D),
        "High": RGBColor(0xB5, 0x54, 0x1F), "Critical": RGBColor(0x7A, 0x1E, 0x1E),
    }.get(result.risk_level, RGBColor(0, 0, 0))
    rp = doc.add_paragraph()
    r1 = rp.add_run(f"Risk Rating: {result.risk_level}  ({result.score}/100)")
    r1.bold = True
    r1.font.color.rgb = risk_color
    doc.add_paragraph(f"CDD tier required: {result.cdd_tier}")
    dec = doc.add_paragraph()
    dec.add_run(f"Decision: {result.decision}").bold = True

    def add_kv_section(title_, kvs):
        doc.add_heading(title_, level=2)
        for k, v in kvs:
            p = doc.add_paragraph()
            p.add_run(f"{k}: ").bold = True
            p.add_run(str(v) if v not in (None, "") else "-")

    add_kv_section("1. Customer Identification", [
        ("Customer type", answers.get("customer_type")),
        ("Full name / legal name", answers.get("full_name")),
        ("DOB / incorporation date", answers.get("dob_or_incorp")),
        ("Nationality / structure", answers.get("nationality") or answers.get("structure_type")),
        ("Country", answers.get("country")),
        ("Occupation / industry", answers.get("occupation") or answers.get("industry")),
        ("ID document", f"{answers.get('id_type')} {answers.get('id_number')}"),
        ("Verification method", answers.get("verification_method")),
    ])

    if answers.get("customer_type") == "Business" and answers.get("ubos"):
        doc.add_heading("2. Beneficial Ownership", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, col in enumerate(["Name", "Role", "Ownership %", "Nationality", "PEP"]):
            hdr[i].text = col
        for u in answers["ubos"]:
            cells = table.add_row().cells
            cells[0].text = str(u.get("name", ""))
            cells[1].text = str(u.get("role", ""))
            cells[2].text = str(u.get("ownership_pct", ""))
            cells[3].text = str(u.get("nationality", ""))
            cells[4].text = "Yes" if u.get("is_pep") else "No"

    add_kv_section("3. Relationship & Activity Profile", [
        ("Purpose of relationship", answers.get("purpose")),
        ("Source of funds", answers.get("source_of_funds")),
        ("Source of wealth", answers.get("source_of_wealth")),
        ("Expected transaction volume", answers.get("expected_volume")),
        ("Expected channels", ", ".join(answers.get("expected_channels", []))),
        ("Expected transacting countries", ", ".join(answers.get("expected_countries", []))),
        ("Cash-intensive business", "Yes" if answers.get("cash_intensive") else "No"),
    ])

    doc.add_heading("4. Screening Results", level=2)
    if result.screening_hits:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, col in enumerate(["Type", "Detail", "Score", "Source"]):
            hdr[i].text = col
        for h in result.screening_hits:
            cells = table.add_row().cells
            cells[0].text = h["type"]
            cells[1].text = h["detail"]
            cells[2].text = str(h["score"])
            cells[3].text = h["source"]
    else:
        doc.add_paragraph("No PEP, sanctions or adverse media matches identified.")

    doc.add_heading("5. Risk Assessment", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, col in enumerate(["Factor", "Points", "Note"]):
        hdr[i].text = col
    for f in result.factors:
        cells = table.add_row().cells
        cells[0].text = f.label
        cells[1].text = f"+{f.points}"
        cells[2].text = f.note

    doc.add_heading("6. Compliance Officer Sign-off", level=2)
    doc.add_paragraph(
        "Prepared by: [Compliance analyst name]      Reviewed by: [MLRO / delegate name]\n"
        "I confirm the identification, verification and risk assessment steps above have been completed "
        "in accordance with the AML/CTF Program.\n"
        "Signature: ______________________      Date: ______________________"
    )

    footer = doc.add_paragraph(
        "Note: This record was generated to accelerate onboarding documentation. All fields must be "
        "verified against original/certified identification documents, and any Critical/High risk rating "
        "or screening hit must be reviewed and approved by the MLRO before the account is activated."
    )
    footer.runs[0].italic = True
    footer.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

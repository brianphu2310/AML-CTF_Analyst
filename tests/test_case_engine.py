"""Case Decision engine, escalation register and SMR Word draft."""
import io
import os
import sys

import pandas as pd
import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core import case as C
from core import escalation as E
from core.model import build_model
from core.ref import AS_OF, MLRO_NAME
from core.smr_doc import build_smr_docx

M = build_model()


def base(**kw):
    c = C.blank_case()
    c.update(client="Test Client Pty Ltd", client_type="Individual", channel="Face-to-face", foreign_tier="Domestic only",
             product="Conveyancing (trust transfer)")
    c.update(kw)
    return c


# ------------------------------------------------------------------ scoring ----
def test_clean_low_risk_case_is_closed():
    r = C.assess(base(), AS_OF)
    assert r["key"] == "close" and not r["escalate"] and not r["clocks"]


def test_score_is_the_sum_of_its_contributions_and_bounded():
    c = base(flags=list(C.RED_FLAGS), amount=900_000)
    assert C.score(c) == 100.0                                   # capped
    c = base(flags=["structuring"], mitigants=["sof_verified"])
    assert C.score(c) == pytest.approx(sum(p for _, p, _ in C.contributions(c)))


@pytest.mark.parametrize("flags,expected", [
    ([], "close"),
    (["structuring"], "edd"),
    (["structuring", "third_party", "reluctant"], "escalate"),
    (["structuring", "third_party", "reluctant", "adverse_media"], "smr"),
])
def test_bands_follow_the_score(flags, expected):
    assert C.assess(base(flags=flags), AS_OF)["key"] == expected


def test_verification_never_raises_the_score_and_clearing_a_flag_never_raises_it():
    c = base(flags=["structuring", "reluctant", "pep"])
    s0 = C.score(c)
    for m in C.MITIGANTS:
        assert C.score({**c, "mitigants": [m]}) <= s0
    for f in c["flags"]:
        assert C.score({**c, "flags": [x for x in c["flags"] if x != f]}) <= s0


def test_hard_stops_override_the_score():
    stop = C.assess(base(hard=["sanctions_confirmed"]), AS_OF)
    assert stop["key"] == "stop" and stop["smr_recommended"] and any("Sanctions Office" in s for s in stop["steps"])
    tf = C.assess(base(hard=["terrorism"]), AS_OF)
    assert tf["key"] == "smr_tf"
    (label, due, note), = [c for c in tf["clocks"] if "terrorism" in c[0]]
    assert due - pd.Timestamp(AS_OF) == pd.Timedelta(hours=24)


def test_smr_clock_is_three_business_days_and_skips_the_weekend():
    thursday = pd.Timestamp("2025-09-18")
    assert C.add_business_days(thursday, 3) == pd.Timestamp("2025-09-23")          # Fri, Mon, Tue
    r = C.assess(base(flags=["structuring", "third_party", "reluctant", "adverse_media"]), thursday)
    assert any(d == pd.Timestamp("2025-09-23") for _, d, _ in r["clocks"])


def test_ttr_only_for_physical_cash_of_10k_or_more():
    assert not C.assess(base(largest_cash=9_999), AS_OF)["ttr"]
    r = C.assess(base(largest_cash=10_000), AS_OF)
    assert r["ttr"] and any("Threshold transaction report" in c[0] for c in r["clocks"])


def test_what_if_flags_the_changes_that_move_the_decision():
    c = base(flags=["structuring", "third_party", "reluctant"])          # escalate band
    base_key = C.assess(c, AS_OF)["key"]
    for kind, label, s2, title2, changes in C.what_if(c, AS_OF):
        alt_key = C.band_for(s2)[1]
        assert changes == (alt_key != base_key)
    assert C.what_if(c, AS_OF)[0][4]                                     # decision-changing rows come first


# ------------------------------------------------------------ from an alert ----
def test_cases_built_from_alerts_are_deterministic_and_consistent_with_their_outcome():
    a = M["alerts"]
    for aid in a["alert_id"].iloc[::7]:
        c1, c2 = C.case_from_alert(M, aid), C.case_from_alert(M, aid)
        assert c1 == c2
        disp = a.set_index("alert_id").loc[aid, "disposition"]
        key = C.assess(c1, AS_OF)["key"]
        if disp == "Escalated to SMR":
            assert key in ("escalate", "smr")
        if disp == "Closed — false positive":
            assert key in ("close", "edd", "escalate") and key != "smr"
        assert abs(c1["amount"] - sum(t[2] for t in c1["transactions"])) < 1e-6


def test_structuring_alerts_are_made_of_deposits_just_under_10k():
    a = M["alerts"]
    aid = a[a["typology"] == "Structuring / smurfing"]["alert_id"].iloc[0]
    tx = C.case_from_alert(M, aid)["transactions"]
    assert tx and all(9_000 <= t[2] < 10_000 for t in tx)


def test_precedent_base_rate_matches_the_alert_ledger():
    hb = C.history_base_rate(M, "Structuring / smurfing", "Medium")
    a = M["alerts"]
    sub = a[(a.typology == "Structuring / smurfing") & (a.risk_tier == "Medium") & ~a.disposition.isin(["Open", "Under investigation"])]
    assert hb["n"] == len(sub) and hb["smr"] == int((sub.disposition == "Escalated to SMR").sum())


# ------------------------------------------------------------ escalations ----
def test_escalation_lifecycle():
    c = base(flags=["structuring", "third_party", "reluctant", "adverse_media"])
    rec = E.new_escalation([], c, AS_OF, "Brian Phu", "  note  ")
    assert rec["status"] == E.STATUS_AWAITING and rec["rationale"] == "note" and rec["mlro"] == MLRO_NAME
    done = E.decide(rec, "suspicion", AS_OF)
    assert done["status"] == E.STATUS_SUSPICION and done["smr_due"] == C.add_business_days(AS_OF, 3)
    assert rec["status"] == E.STATUS_AWAITING                            # original untouched
    assert E.decide(rec, "close", AS_OF)["smr_due"] is None


def test_escalation_history_reconciles_with_the_smr_register():
    h = E.history(M)
    assert len(h) == len(M["smrs"]) and set(h["smr_id"]) == set(M["smrs"]["smr_id"])
    reg = E.register(M, [E.new_escalation([], base(), AS_OF, "Brian Phu")])
    assert len(reg) == len(h) + 1 and reg.iloc[0]["source"] == "This session"


# ------------------------------------------------------------ SMR Word draft ----
def _text(docx_bytes):
    d = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [cell.text for cell in row.cells]
    parts += [p.text for p in d.sections[0].header.paragraphs] + [p.text for p in d.sections[0].footer.paragraphs]
    return "\n".join(parts)


def test_smr_docx_contains_every_section_and_the_case_facts():
    s = M["smrs"].iloc[-1]
    c = C.case_from_alert(M, s["alert_id"])
    txt = _text(build_smr_docx(c, s["escalated"], smr_ref=s["smr_id"], escalated_on=s["escalated"], rationale="Analyst note here."))
    for heading in ["Report details", "Person or organisation the suspicion relates to", "Transactions", "Grounds for suspicion",
                    "Indicators and verification", "Action taken by the reporting entity", "Review and sign-off"]:
        assert heading in txt
    assert c["client"] in txt and s["smr_id"] in txt and s["alert_id"] in txt and "Analyst note here." in txt
    assert MLRO_NAME in txt and "s123" in txt and "Synthetic" in txt
    assert f"AUD {sum(t[2] for t in c['transactions']):,.0f}" in txt


def test_smr_docx_flags_terrorism_financing_deadline_and_ttr():
    c = base(hard=["terrorism"], largest_cash=15_000, transactions=[(pd.Timestamp("2025-09-10"), "Cash deposit to trust", 15_000.0, "In", "Client")])
    txt = _text(build_smr_docx(c, AS_OF))
    assert "24 hours" in txt and "Terrorism financing" in txt and "threshold transaction report" in txt.lower()
